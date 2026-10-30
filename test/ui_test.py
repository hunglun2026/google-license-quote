"""實際在瀏覽器操作網頁，核對算式、匯出 xlsx、以及抓不到牌告時的行為。"""
import sys, tempfile, os
from playwright.sync_api import sync_playwright

URL = "http://localhost:8788/"
fails = []

def check(name, ok, detail=""):
    print(("OK   " if ok else "FAIL ") + name, detail)
    if not ok:
        fails.append(name)

def n(s):
    return float(s.replace(",", "").split("（")[0])

with sync_playwright() as p:
    b = p.chromium.launch(headless=True)
    pg = b.new_page(accept_downloads=True)
    errs = []
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(URL)
    pg.wait_for_function("document.getElementById('sell').textContent !== '--'")
    sell = float(pg.inner_text("#sell"))
    check("狀態顯示", pg.inner_text("#status") in ("即時",) or "未更新" in pg.inner_text("#status"), pg.inner_text("#status"))

    pg.fill("#unit", "100"); pg.fill("#qty", "1000"); pg.fill("#buffer", "0.3")
    pg.fill("#fee", "500"); pg.fill("#markup", "10"); pg.check("#tax")
    usd = 100 * 1000
    fx = usd * (sell + 0.3); cost = fx + 500; net = cost * 1.1; quote = net * 1.05
    got = {k: n(pg.inner_text("#" + k)) for k in ("rUsd", "rFx", "rCost", "rNet", "rTax", "rQuote", "rPer")}
    check("美金總額", abs(got["rUsd"] - usd) < 0.01, got["rUsd"])
    check("含緩衝成本", abs(got["rFx"] - round(fx)) <= 1, f'{got["rFx"]} vs {fx:.1f}')
    check("成本合計", abs(got["rCost"] - round(cost)) <= 1, f'{got["rCost"]} vs {cost:.1f}')
    check("未稅報價", abs(got["rNet"] - round(net)) <= 1, f'{got["rNet"]} vs {net:.1f}')
    check("營業稅", abs(got["rTax"] - round(net * 0.05)) <= 1)
    check("最終報價", abs(got["rQuote"] - round(quote)) <= 1, f'{got["rQuote"]} vs {quote:.1f}')
    check("每單位", abs(got["rPer"] - round(quote / 1000)) <= 1)

    # 手動議價匯率
    pg.select_option("#mode", "manual"); pg.fill("#manual", "32")
    pg.uncheck("#tax"); pg.fill("#buffer", "0"); pg.fill("#fee", "0"); pg.fill("#markup", "0")
    check("手動匯率成本", abs(n(pg.inner_text("#rCost")) - 3200000) <= 1, pg.inner_text("#rCost"))
    pg.select_option("#mode", "auto")

    # 匯出 xlsx
    with pg.expect_download() as d:
        pg.click("#exportBtn")
    path = os.path.join(tempfile.gettempdir(), "gq_test.xlsx")
    d.value.save_as(path)
    from openpyxl import load_workbook
    ws = load_workbook(path).active
    vals = {r[0].value: r[1].value for r in ws.iter_rows() if r[0].value}
    check("xlsx 有銀行時間", bool(vals.get("銀行公告時間")), vals.get("銀行公告時間"))
    check("xlsx 採用匯率", abs(float(vals["採用匯率（台幣/美金）"]) - sell) < 0.0005, vals["採用匯率（台幣/美金）"])

    # 抓不到牌告：不可沿用舊數字
    pg2 = b.new_page()
    pg2.route("**/api/rate", lambda r: r.fulfill(status=502, content_type="application/json", body='{"ok":false,"error":"測試斷線"}'))
    pg2.goto(URL)
    pg2.wait_for_function("document.getElementById('status').textContent !== '讀取中'")
    pg2.fill("#unit", "100"); pg2.fill("#qty", "10")
    check("斷線顯示錯誤", pg2.inner_text("#status") == "無法取得牌告", pg2.inner_text("#status"))
    check("斷線不算成本", pg2.inner_text("#rCost") == "--", pg2.inner_text("#rCost"))
    check("斷線不能匯出", pg2.is_disabled("#exportBtn"))

    # 過期資料
    pg3 = b.new_page()
    pg3.route("**/api/rate", lambda r: r.fulfill(content_type="application/json",
        body='{"ok":true,"bankTime":"2026/10/29 09:00:00","spotBuy":32.6,"spotSell":32.7,"fetchedAt":"x"}'))
    pg3.goto(URL)
    pg3.wait_for_function("document.getElementById('status').textContent !== '讀取中'")
    check("過期提示", "未更新" in pg3.inner_text("#status"), pg3.inner_text("#status"))

    pg.screenshot(path=os.path.join(tempfile.gettempdir(), "gq_shot.png"), full_page=True)
    check("無 JS 錯誤", not errs, errs)
    b.close()

print("失敗:", fails)
sys.exit(1 if fails else 0)
